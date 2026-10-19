import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r"C:\Users\XOSIYAT\Desktop\kovsh"
OUT = os.path.join(BASE, "output", "revised")
REPORT = "Natijalar_jamlanmasi_REVIDERLANGAN.docx"

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('YOLO EKSPERIMENTI — REAL NATIJALAR JAMLANMASI')
r.bold = True
r.font.size = Pt(15)
r.font.name = 'Times New Roman'

def para(text, bold=False, size=12, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = 'Times New Roman'
    return p

# ---- Experiment meta ----
para('Tajriba sharoiti:', bold=True)
meta = [
    'Video: static/video.mp4  |  Kadrlar: 200  |  Deteksiyalar: 920',
    'Model: exp-2.pt  |  conf >= 0.25  |  imgsz = 320  |  Video o\u2018lchami 2592 \u00d7 1944',
]
for m in meta:
    para(m)

# ---- Per-class summary table ----
doc.add_page_break()
para('1. Sinflar bo\u2018yicha tavsifiy statistika', bold=True)
t = doc.add_table(rows=11, cols=4)
t.style = 'Table Grid'
headers = ['Ko\u2018rsatkich', 'bucket-teeth', 'truck', 'Izoh']
perclass = [
    ['Deteksiyalar soni', '820', '100', '920 dan'],
    ['Ulush', '89.1%', '10.9%', 'Detection distribution'],
    ['O\u2018rtacha confidence', '0.472', '0.623', ''], 
    ['Standart chetlanish', '\u00b10.156', '\u00b10.179', ''],
    ['Confidence diapazoni', '0.250\u20130.859', '0.260\u20130.855', 'conf \u2265 0.25'],
    ['Bbox kenglik (ort.)', '85 px', '2514 px', ''],
    ['Bbox balandlik (ort.)', '224 px', '1648 px', ''],
    ['Maydon (ort.)', '19 100 px\u00b2', '4 139 412 px\u00b2', ''],
    ['Maydon (median)', '18 622 px\u00b2', '4 334 854 px\u00b2', ''],
    ['Maydon diapazoni', '10 758 \u2013 30 520', '3 333 832 \u2013 4 648 488', ''],
]
for i, row in enumerate(perclass):
    cells = t.rows[i+1].cells
    for j, val in enumerate(row):
        c = cells[j]
        c.text = ''
        rr = c.paragraphs[0].add_run(val)
        rr.font.name = 'Times New Roman'
        rr.font.size = Pt(10)
        if j == 0:
            rr.bold = True
for j, htext in enumerate(headers):
    c = t.rows[0].cells[j]
    c.text = ''
    rr = c.paragraphs[0].add_run(htext)
    rr.bold = True
    rr.font.name = 'Times New Roman'
    rr.font.size = Pt(10)

# ---- Retention table ----
para('2. Retention tahlili (haqiqiy Precision/Recall EMAS!)', bold=True)
para('Confidence chegarasi oshganda saqlanib qolgan deteksiyalar ulushi:', italic=False, size=11)
t2 = doc.add_table(rows=10, cols=3)
t2.style = 'Table Grid'
h2 = ['Chegara', 'bucket-teeth (%)', 'truck (%)']
for j, htext in enumerate(h2):
    c = t2.rows[0].cells[j]
    c.text = ''
    rr = c.paragraphs[0].add_run(htext)
    rr.bold = True
    rr.font.name = 'Times New Roman'
    rr.font.size = Pt(10)
retention = [
    ['\u2265 0.25', '100', '100'],
    ['\u2265 0.30', '85', '96'],
    ['\u2265 0.40', '61', '81'],
    ['\u2265 0.50', '41', '73'],
    ['\u2265 0.60', '23', '62'],
    ['\u2265 0.70', '11', '51'],
    ['\u2265 0.80', '2', '11'],
    ['\u2265 0.85', '0', '2'],
    ['\u2265 0.90', '0', '0'],
]
for i, row in enumerate(retention):
    cells = t2.rows[i+1].cells
    for j, val in enumerate(row):
        c = cells[j]
        c.text = ''
        rr = c.paragraphs[0].add_run(val)
        rr.font.name = 'Times New Roman'
        rr.font.size = Pt(10)

# ---- FPS ----
para('3. Performance (benchmark)', bold=True)
para('CPU, imgsz = 640:', bold=False)
fps = [
    'O\u2018rtacha inference: 64.3 ms/kadr',
    'FPS: 15.6',
    'Median: 60.0 ms | p95: 93.8 ms | min: 53.2 ms | max: 126.9 ms',
    'Video tezligi: 20 fps',
    'GPU: o\u2018lchash kutilmoqda',
]
for f in fps:
    para(f, size=11)

doc.add_page_break()
para('4. Grafiklar', bold=True)
graphics = [
    ('1_confidence_distribution.png', 'Confidence taqsimoti'),
    ('2_class_distribution.png', 'Deteksiyalarning sinflar bo\u2018yicha ulushi'),
    ('3_retention_curve.png', 'Retention egri'),
    ('4_f1_retention.png', 'Retention asosidagi F1'),
    ('5_detections_per_frame.png', 'Har kadrda deteksiyalar soni'),
    ('6_box_size_distribution.png', 'Obyekt o\u2018lchami taqsimoti'),
    ('7_class_confidence.png', 'O\u2018rtacha confidence \u00b1 std'),
]
for fname, cap in graphics:
    path = os.path.join(OUT, fname)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Cm(15))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = cp.add_run(cap)
        rr.italic = True
        rr.font.name = 'Times New Roman'
        rr.font.size = Pt(9)

# ---- Important note ----
para('5. Maqola yozishda muhim eslatmalar', bold=True)
notes = [
    '3- va 4-grafiklar "Precision/Recall/F1" emas, balki retention egridir. Sarlavhada shuningdek yozilishi kerak.',
    'Haqiqiy Ground-Truth asosidagi Precision/Recall/F1/mAP hali hisoblanmagan (annotatsiya talab qilinadi). Maqolada buni "kelgusi ish" sifatida ko\u2018rsatish lozim.',
    'FPS: CPU uchun 15.6 FPS qayd etiladi; GPU o\u2018lchovi olinishi tavsiya etiladi.',
    'Retention raqamlari: bucket-teeth uchun 0.30 chegara 85%, truck uchun 96% saqlanishni beradi.',
]
for i, n in enumerate(notes, 1):
    para(f'{i}. {n}', size=11)

out_path = os.path.join(OUT, REPORT)
doc.save(out_path)
print('SAVED:', out_path)